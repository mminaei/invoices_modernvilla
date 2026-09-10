import streamlit as st
import docx
import io
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree
from datetime import datetime

BASE_DIR = Path(__file__).parent
PRICE_FILE = BASE_DIR / 'Villa_Design_Prices.xlsx'
WORD_TEMPLATE = BASE_DIR / 'پروپوزال طراحی فول پکیج.docx'


def _to_number(value):
    digits = str(value).translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹', '0123456789'))
    return float(digits.replace(',', '').strip())


def calculate_prices(bana, land, interior):
    prices = {
        'plan': 350, 'land': 100, 'phase2': 100,
        'struct': 70, 'mech': 70, 'elec': 70,
        'interior': 400, 'interior_p2': 80
    }
    if not PRICE_FILE.exists():
        return prices

    namespace = {'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(PRICE_FILE) as workbook:
        sheet = ElementTree.fromstring(workbook.read('xl/worksheets/sheet1.xml'))
    rows = sheet.findall('.//main:row', namespace)
    for row in rows[1:]:
        cells = row.findall('main:c', namespace)
        values = [''.join(cell.itertext()).strip() for cell in cells]
        if len(values) < 6 or not values[0]:
            continue
        price_values = [_to_number(value) for value in values[1:6]]
        area_numbers = [_to_number(part) for part in values[0].split('-')]
        if len(area_numbers) == 1 and bana >= area_numbers[0]:
            selected = True
        elif len(area_numbers) == 2 and area_numbers[0] <= bana <= area_numbers[1]:
            selected = True
        else:
            selected = False
        if selected:
            prices.update(plan=price_values[0], phase2=price_values[1], struct=price_values[2],
                          elec=price_values[3], mech=price_values[4])
            break
    return prices

def _format_amount(million_toman):
    rial = f'{million_toman * 10_000_000:,.0f}'
    toman = f'{million_toman:,.0f}'
    return f'{rial} ریال\n({toman} میلیون تومان)'


def _replace_template_content(doc, client, title, project, bana, land, interior, location, prices):
    for paragraph in doc.paragraphs:
        text = paragraph.text
        if text == 'سرکار خانم …':
            paragraph.text = f'سرکار خانم {client}' if title == 'خانم' else ''
        elif text == 'جناب آقای ...':
            paragraph.text = f'جناب آقای {client}' if title == 'آقا' else ''
        elif text.startswith('با توجه به درخواست'):
            paragraph.text = (
                f'با توجه به درخواست مطرح شده در خصوص طراحی {project} به مساحت '
                f'{bana} مترمربع در زمین به مساحت حدود {land} مترمربع واقع در '
                f'{location}، با متراژ طراحی داخلی {interior} مترمربع، '
                'شرح خدمات به شما ارائه می گردد.'
            )
        elif text.startswith(('1-', '2-', '3-', '4-', '5-', '6-')):
            payment_prices = {
                '1-': prices['plan'], '2-': prices['phase2'],
                '3-': prices['struct'], '4-': prices['mech'],
                '5-': prices['interior'], '6-': prices['interior_p2']
            }
            payment_key = text[:2]
            amount = _format_amount(payment_prices[payment_key])
            paragraph.text = text.replace('0.0.00.000.000', amount).replace(
                '0.000.000.000', amount
            )

    service_prices = [
        prices['plan'], prices['plan'], prices['land'], prices['phase2'],
        prices['struct'], prices['mech'], prices['mech'], prices['elec'],
        prices['interior'], prices['interior_p2']
    ]
    price_index = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if '0.000.000.000' in cell.text:
                    cell.text = _format_amount(service_prices[price_index])
                    price_index += 1

def generate_documents(client, title, project, bana, land, interior, location):
    prices = calculate_prices(bana, land, interior)
    
    # 1. تولید PDF در حافظه (بدون ذخیره روی هارد)
    html_template = f"""
    <html lang="fa" dir="rtl">
    <head><style>body {{ font-family: Tahoma; }}</style></head>
    <body>
        <h2>به نام خدا</h2>
        <p>{'سرکار خانم' if title == 'خانم' else 'جناب آقای'} <strong>{client}</strong></p>
        <p>پروپوزال طراحی {project} به مساحت {bana} متر مربع واقع در {location}.</p>
    </body>
    </html>
    """
    pdf_buffer = None
    try:
        from weasyprint import HTML
        pdf_buffer = io.BytesIO()
        HTML(string=html_template).write_pdf(pdf_buffer)
        pdf_buffer.seek(0)
    except (ImportError, OSError):
        pass
    
    # 2. تولید Word در حافظه
    doc = docx.Document(str(WORD_TEMPLATE)) if WORD_TEMPLATE.exists() else docx.Document()
    _replace_template_content(doc, client, title, project, bana, land, interior, location, prices)
    
    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)
    
    return pdf_buffer, docx_buffer

# --- رابط کاربری ---
st.set_page_config(page_title="سیستم صدور پروپوزال", page_icon="📄")
st.markdown("<h2 style='text-align: center;'>سیستم هوشمند صدور پروپوزال</h2>", unsafe_allow_html=True)

with st.form("proposal_form"):
    client_input = st.text_input('نام کارفرما')
    title_input = st.selectbox('عنوان مخاطب', ['خانم', 'آقا'])
    project_input = st.text_input('عنوان پروژه (مثل: ویلای تریبلکس)')
    bana_input = st.number_input('متراژ بنا (مترمربع)', min_value=0, step=10)
    land_input = st.number_input('متراژ زمین/محوطه (مترمربع)', min_value=0, step=10)
    interior_input = st.number_input('متراژ طراحی داخلی', min_value=0, step=10)
    location_input = st.text_input('محل پروژه')
    
    submitted = st.form_submit_button("محاسبه و تولید فایل‌ها")

if submitted:
    if not client_input or not bana_input:
        st.error('لطفاً اطلاعات ضروری را وارد کنید!')
    else:
        with st.spinner('در حال پردازش و تولید فایل‌ها...'):
            pdf_buf, docx_buf = generate_documents(
                client_input, title_input, project_input, bana_input,
                land_input, interior_input, location_input
            )
            
            st.success('✅ فایل‌ها با موفقیت تولید شدند!')
            
            col1, col2 = st.columns(2)
            with col1:
                if pdf_buf:
                    st.download_button(
                        label="📥 دانلود PDF",
                        data=pdf_buf,
                        file_name=f"Proposal_{client_input}.pdf",
                        mime="application/pdf",
                        use_container_width=True
                    )
                else:
                    st.warning('PDF در این محیط در دسترس نیست؛ وابستگی‌های سیستمی WeasyPrint نصب نشده‌اند.')
            with col2:
                st.download_button(
                    label="📥 دانلود Word",
                    data=docx_buf,
                    file_name=f"Proposal_{client_input}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )